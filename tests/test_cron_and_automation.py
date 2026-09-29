"""
Tests - CORD Persistent Task Automation, Cron Scheduler & Document Generator
"""

import os
import time
import pytest
from pathlib import Path
import docx

from cord.cron.job_model import CronJob
from cord.cron.schedule_parser import ScheduleParser
from cord.cron.cron_manager import CronManager
from cord.cron.report_generator import ReportGenerator
from cord.cron.email_notifier import EmailNotifier
from cord.tools.cron.schedule_cron_tool import ScheduleCronJobTool
from cord.tools.cron.generate_report_tool import GenerateReportTool
from cord.tools.cron.send_email_tool import SendEmailTool


def test_schedule_parser_expressions():
    parser = ScheduleParser()
    now = time.time()

    # 1. Relative interval: every 24h
    next_24h = parser.calculate_next_run("every 24h", after_ts=now)
    assert 86390 <= (next_24h - now) <= 86410

    # 2. Relative interval: every 30m
    next_30m = parser.calculate_next_run("every 30m", after_ts=now)
    assert 1790 <= (next_30m - now) <= 1810

    # 3. Specific time of day: every morning at 8:00
    next_8am = parser.calculate_next_run("every morning at 8:00", after_ts=now)
    assert next_8am > now

    # 4. Standard cron: 0 8 * * *
    next_cron = parser.calculate_next_run("0 8 * * *", after_ts=now)
    assert next_cron > now


def test_cron_manager_lifecycle(tmp_path):
    jobs_file = tmp_path / "cron_jobs.json"
    mgr = CronManager(jobs_file=jobs_file)

    job1 = mgr.add_job(
        name="Daily AI Briefing",
        schedule_expr="every morning at 8:00",
        prompt="Gather top AI news from the past 24 hours",
        action_type="generate_report",
        recipient_email="test@example.com",
    )

    assert job1.id in mgr.jobs
    assert job1.name == "Daily AI Briefing"
    assert job1.next_run > time.time()

    # Persistence verification
    mgr2 = CronManager(jobs_file=jobs_file)
    assert len(mgr2.list_jobs()) == 1
    assert mgr2.get_job(job1.id).action_type == "generate_report"

    # Due jobs
    job1.next_run = time.time() - 10
    mgr2.jobs[job1.id].next_run = time.time() - 10
    due = mgr2.get_due_jobs()
    assert len(due) == 1

    mgr2.remove_job(job1.id)
    assert len(mgr2.list_jobs()) == 0


def test_docx_word_report_generation(tmp_path):
    docx_path = tmp_path / "briefing.docx"
    sections = [
        {"heading": "1. Overview", "content": "Global AI ecosystem report for 2026."},
        {"heading": "2. Top Breakthroughs", "bullets": ["Open-source reasoning models", "Autonomous desktop agents"]},
        {
            "heading": "3. Benchmarks",
            "table": {
                "headers": ["Model", "Reasoning Score", "Latency"],
                "rows": [["Claude 3.7", "98.2", "450ms"], ["DeepSeek R1", "97.8", "620ms"]],
            },
        },
    ]

    saved_file = ReportGenerator.generate_docx(
        title="CORD AI Research Digest",
        subtitle="Automated intelligence briefing",
        sections=sections,
        output_path=str(docx_path),
    )

    assert os.path.exists(saved_file)
    assert os.path.getsize(saved_file) > 1000

    # Read back with python-docx
    doc = docx.Document(saved_file)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "CORD AI Research Digest" in full_text
    assert "Global AI ecosystem report" in full_text
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[0].cells[0].text == "Model"


def test_email_notifier_outbox_queue(tmp_path):
    cfg_path = tmp_path / "email.json"
    notifier = EmailNotifier(config_path=cfg_path)
    # Set outbox to tmp_path
    notifier.outbox_dir = tmp_path / "outbox"
    notifier.outbox_dir.mkdir(parents=True, exist_ok=True)

    dummy_doc = tmp_path / "summary.docx"
    dummy_doc.write_text("dummy doc content", encoding="utf-8")

    res = notifier.send(
        to_email="user@company.com",
        subject="Daily Briefing Ready",
        body_text="Here is your requested briefing.",
        attachments=[str(dummy_doc)],
    )

    assert res["success"] is True
    assert res["mode"] in ("outbox_queued", "outbox_fallback")
    assert os.path.exists(res["outbox_path"])


@pytest.mark.asyncio
async def test_cron_tools_execution(tmp_path):
    # Test GenerateReportTool
    rep_tool = GenerateReportTool()
    out_docx = tmp_path / "tool_report.docx"
    res_tool = await rep_tool.execute(
        title="Weekly Analytics",
        output_path=str(out_docx),
        sections=[{"heading": "Status", "content": "Everything optimal."}],
    )
    assert res_tool.success
    assert os.path.exists(out_docx)

    # Test ScheduleCronJobTool
    sched_tool = ScheduleCronJobTool()
    sched_res = await sched_tool.execute(
        name="Sync Media",
        schedule="every 12h",
        prompt="Backup media to nas-01",
        target_node="nas-01",
    )
    assert sched_res.success
    assert "Sync Media" in sched_res.output
